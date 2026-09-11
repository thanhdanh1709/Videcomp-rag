"""Trình phân tích và băm nhỏ tệp tải lên (PDF, Word DOCX, TXT, MD) cho Videcomp-rag."""
from __future__ import annotations

import hashlib
import io
import logging
import re
from typing import BinaryIO, Callable

from ..schemas.evidence import EvidenceCandidate

logger = logging.getLogger(__name__)


def _clean_text(text: str) -> str:
    # Chuẩn hóa khoảng trắng và dòng trống
    text = re.sub(r"\r\n|\r", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_file(
    filename: str,
    content: bytes,
    progress_callback: Callable[[float, str], None] | None = None,
) -> list[tuple[str, str]]:
    """Trích xuất văn bản từ tệp dưới dạng danh sách [(tiêu_đề_phần, nội_dung)]."""
    ext = filename.lower().split(".")[-1]
    sections: list[tuple[str, str]] = []

    if progress_callback:
        progress_callback(5.0, f"Bắt đầu đọc tệp {filename}...")

    if ext == "pdf":
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(content))
        total_pages = len(reader.pages)
        for page_idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = _clean_text(text)
            if text:
                sections.append((f"Trang {page_idx}", text))
            if progress_callback and total_pages > 0:
                pct = 5.0 + (page_idx / total_pages) * 55.0
                progress_callback(pct, f"Đang trích xuất trang {page_idx}/{total_pages} ({int(pct)}%)...")

    elif ext in ("docx", "doc"):
        if progress_callback:
            progress_callback(15.0, "Đang phân tích cấu trúc văn bản và bảng biểu Word...")
        import docx

        doc = docx.Document(io.BytesIO(content))
        paragraphs = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                paragraphs.append(t)
        # Bao gồm cả bảng biểu nếu có
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        if progress_callback:
            progress_callback(45.0, "Đang tổng hợp nội dung văn bản Word...")
        full_doc_text = _clean_text("\n".join(paragraphs))
        if full_doc_text:
            sections.append(("Nội dung văn bản", full_doc_text))

    else:  # txt, md, csv, html, v.v.
        if progress_callback:
            progress_callback(20.0, "Đang giải mã và làm sạch văn bản...")
        try:
            raw_text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                raw_text = content.decode("utf-16")
            except Exception:
                raw_text = content.decode("latin-1", errors="ignore")

        clean = _clean_text(raw_text)
        if clean:
            sections.append(("Toàn văn tài liệu", clean))

    if progress_callback:
        progress_callback(60.0, f"Đã trích xuất xong {len(sections)} phần nội dung.")

    return sections


def parse_file_to_candidates(
    filename: str,
    content: bytes,
    chunk_size: int = 400,
    chunk_overlap: int = 50,
    progress_callback: Callable[[float, str], None] | None = None,
) -> list[EvidenceCandidate]:
    """Phân rã tệp tải lên thành các đoạn EvidenceCandidate có cấu trúc chuẩn."""
    sections = extract_text_from_file(filename, content, progress_callback=progress_callback)
    candidates: list[EvidenceCandidate] = []

    clean_doc_name = re.sub(r"[^\w\s.-]", "", filename).strip()
    file_hash = hashlib.md5(filename.encode("utf-8", errors="ignore")).hexdigest()[:6]
    doc_id = f"FILE_{file_hash}_{clean_doc_name[:25]}"

    if progress_callback:
        progress_callback(65.0, "Bắt đầu băm nhỏ văn bản (chunking) theo cửa sổ trượt...")

    chunk_seq = 1
    total_sections = max(1, len(sections))
    for sec_idx, (sec_title, sec_text) in enumerate(sections, start=1):
        words = sec_text.split()
        if not words:
            continue

        start = 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            chunk_words = words[start:end]
            chunk_str = " ".join(chunk_words)

            chunk_id = f"up_{file_hash}_{chunk_seq}"
            candidates.append(
                EvidenceCandidate(
                    chunk_id=chunk_id,
                    doc_id=doc_id,
                    text=f"【{filename} - {sec_title}】: {chunk_str}",
                    source_url=f"attachment://{filename}",
                    parent_path=[filename, sec_title, f"Đoạn {chunk_seq}"],
                    sparse_rank=chunk_seq,
                    dense_rank=chunk_seq,
                    rrf_score=0.95 - (chunk_seq * 0.01),
                )
            )
            chunk_seq += 1
            if end >= len(words):
                break
            start += chunk_size - chunk_overlap

        if progress_callback:
            pct = 65.0 + (sec_idx / total_sections) * 30.0
            progress_callback(pct, f"Đang chia đoạn {sec_title} (đoạn {chunk_seq - 1}, {int(pct)}%)...")

    if progress_callback:
        progress_callback(96.0, f"Đã tạo {len(candidates)} đoạn trích xuất. Đang lưu trữ...")

    return candidates
